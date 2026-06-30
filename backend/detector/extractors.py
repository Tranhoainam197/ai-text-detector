"""
Extract plain text from uploaded files.
Supported: .txt, .docx, .pdf
"""
import io


def extract_txt(data: bytes) -> str:
    """Decode plain text. Try UTF-8 first, fall back to latin-1."""
    try:
        return data.decode("utf-8")
    except UnicodeDecodeError:
        return data.decode("latin-1", errors="replace")


def extract_docx(data: bytes) -> str:
    """Pull paragraphs out of a Word document."""
    from docx import Document
    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    # Also pull text from tables, if any
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    parts.append(cell.text)
    return "\n\n".join(parts)


def extract_pdf(data: bytes) -> str:
    """Pull text from a PDF, page by page."""
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(data))
    parts = []
    for page in reader.pages:
        text = page.extract_text() or ""
        if text.strip():
            parts.append(text)
    return "\n\n".join(parts)


def extract(filename: str, data: bytes) -> str:
    """
    Dispatch to the right extractor based on file extension.
    Raises ValueError for unsupported types.
    """
    name = filename.lower()
    if name.endswith(".txt"):
        return extract_txt(data)
    if name.endswith(".docx"):
        return extract_docx(data)
    if name.endswith(".pdf"):
        return extract_pdf(data)
    raise ValueError(f"Unsupported file type: {filename}. Supported: .txt, .docx, .pdf")